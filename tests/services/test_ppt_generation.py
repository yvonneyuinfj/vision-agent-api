"""PPT 生成链路的本机冒烟：资料解析 / 大纲降级 / pptx 渲染可回读。

只测纯逻辑与 python-pptx 往返，不依赖 DB/LLM/服务器；`uv run pytest` 即可跑通，
用来证明"生成→可下载文件"的产物本身是合法 .pptx。
"""
from __future__ import annotations

import io

import pytest

from app.services.doc_parse import extract_text
from app.services.ppt_render import render_pptx
from app.services.ppt_service import (
    PptService, _normalize_outline, _outline_from_template, _parse_outline,
)


def test_extract_text_plain_and_unsupported():
    text, err = extract_text("notes.md", "# 标题\n内容一\n内容二".encode("utf-8"))
    assert err == "" and "内容一" in text

    text, err = extract_text("x.xyz", b"whatever")
    assert text == "" and "不支持" in err


def test_extract_text_docx():
    from docx import Document

    doc = Document()
    doc.add_paragraph("段落一")
    doc.add_paragraph("段落二")
    buf = io.BytesIO()
    doc.save(buf)
    text, err = extract_text("m.docx", buf.getvalue())
    assert err == "" and "段落一" in text and "段落二" in text


def test_parse_outline_accepts_and_rejects():
    good = '```json\n{"title":"T","subtitle":"S","sections":[{"heading":"H","bullets":["a","b"]}],"closing":"END"}\n```'
    outline = _parse_outline(good)
    assert outline and outline["title"] == "T" and outline["sections"][0]["bullets"] == ["a", "b"]

    assert _parse_outline("not json at all") is None
    assert _parse_outline('{"title":"x"}') is None  # 无 sections 触发降级


def test_template_outline_from_params():
    outline = _outline_from_template({"topic": "季度总结", "audience": "评审组", "request": "目标A\n目标B"}, "")
    assert outline["title"] == "季度总结"
    assert any("目标A" in " ".join(s["bullets"]) for s in outline["sections"])


@pytest.mark.asyncio
async def test_service_falls_back_without_llm():
    svc = PptService(llm=None, file_repo=None, task_repo=None)
    outline, engine = await svc.build_outline({"topic": "无模型主题", "request": "要点"}, "")
    assert engine == "template" and outline["sections"]


def test_render_produces_valid_pptx():
    from pptx import Presentation

    outline = {
        "title": "演示标题",
        "subtitle": "副标题",
        "sections": [{"heading": "第一节", "bullets": ["要点1", "要点2"]}, {"heading": "第二节", "bullets": ["要点3"]}],
        "closing": "谢谢",
    }
    blob = render_pptx(outline, "科技汇报")
    assert isinstance(blob, bytes) and len(blob) > 1000

    prs = Presentation(io.BytesIO(blob))  # 能被 python-pptx 重新打开 = 合法文件
    assert len(prs.slides._sldIdLst) == 1 + 2 + 1  # 封面 + 两节 + 封底
    first = prs.slides[0].shapes.title.text
    assert first == "演示标题"


def test_normalize_manual_outline():
    norm = _normalize_outline({
        "title": "T", "subtitle": "S",
        "sections": [{"heading": "H", "bullets": "a\nb"}, {"heading": "H2", "bullets": ["c"]}],
    })
    assert norm and norm["title"] == "T"
    assert norm["sections"][0]["bullets"] == ["a", "b"]  # 字符串按行拆分
    assert norm["sections"][1]["bullets"] == ["c"]
    assert _normalize_outline({"sections": []}) is None
    assert _normalize_outline("nope") is None


@pytest.mark.asyncio
async def test_service_prefers_manual_outline():
    svc = PptService(llm=None, file_repo=None, task_repo=None)
    params = {"topic": "x", "outline": {"title": "我的", "sections": [{"heading": "A", "bullets": ["1", "2"]}]}}
    outline, engine = await svc.build_outline(params, "")
    assert engine == "manual" and outline["title"] == "我的" and outline["sections"][0]["heading"] == "A"


def test_render_with_template_clears_example_slides():
    from pptx import Presentation

    # 造一个含 3 页示例的“模板”
    tpl = Presentation()
    for _ in range(3):
        tpl.slides.add_slide(tpl.slide_layouts[1])
    tbuf = io.BytesIO()
    tpl.save(tbuf)
    template_bytes = tbuf.getvalue()

    outline = {
        "title": "套模板", "subtitle": "s",
        "sections": [{"heading": "一", "bullets": ["x", "y"]}], "closing": "END",
    }
    blob = render_pptx(outline, template_bytes=template_bytes)
    prs = Presentation(io.BytesIO(blob))
    # 模板自带示例页应被清空，只剩 封面 + 1 节 + 封底 = 3
    assert len(prs.slides._sldIdLst) == 3
    assert prs.slides[0].shapes.title.text == "套模板"
