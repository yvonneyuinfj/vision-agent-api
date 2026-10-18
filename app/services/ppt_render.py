"""PPTX 渲染：把结构化大纲（dict）画成 .pptx 字节。

纯函数、无外部副作用，便于单测冒烟（本机装 python-pptx 即可跑通）。
大纲契约（LLM / 模板 / 手动三条产出共用）：
    {
      "title": str, "subtitle": str,
      "sections": [ {"heading": str, "bullets": [str, ...]}, ... ],
      "closing": str            # 可选，末页致谢/口号；缺省用通用文案
    }

两种渲染模式：
- 无模板（template_bytes=None）：用 python-pptx 默认版式，按 style 上强调色/字体；
- 上传模板（template_bytes 为 .pptx 字节）：载入模板（保留其母版/主题/背景/Logo），
  清空模板自带示例页，按名称/占位符启发式挑选标题版式与内容版式，把大纲填进占位符；
  版式缺对应占位符时回退为手工文本框，保证内容一定可见。模板模式不覆盖字体，交由主题决定。

渲染只在打开端用查看器字体，服务器不需要装字体。
"""
from __future__ import annotations

import io
from datetime import datetime, timezone, timedelta
from typing import Any

_BEIJING = timezone(timedelta(hours=8))

# 风格 → 标题强调色（RGB hex）；未识别用默认蓝（仅无模板模式使用）
_STYLE_ACCENT: dict[str, int] = {
    "简洁商务": 0x2368E4,
    "科技汇报": 0x0E7C86,
    "学术交流": 0x7A3E9B,
}
_DEFAULT_ACCENT = 0x2368E4
# 正文里希望使用的中文字体名（打开端缺失时自动回退）
_BODY_FONT = "Microsoft YaHei"

# 每页要点上限，避免溢出
_MAX_BULLETS = 8


def _rgb(hex_int: int) -> tuple[int, int, int]:
    return (hex_int >> 16) & 255, (hex_int >> 8) & 255, hex_int & 255


def _rgb_obj(triple: tuple[int, int, int]):
    from pptx.dml.color import RGBColor

    return RGBColor(*triple)


def _set_font(run: Any, *, size: int | None = None, bold: bool | None = None, color: tuple[int, int, int] | None = None) -> None:
    from pptx.util import Pt

    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    run.font.name = _BODY_FONT
    if color is not None:
        run.font.color.rgb = _rgb_obj(color)


def _style_title(shape: Any, accent: tuple[int, int, int], size: int) -> None:
    """给占位符标题上色加粗（runs 为空时跳过，不抛）。"""
    para = shape.text_frame.paragraphs[0]
    if para.runs:
        _set_font(para.runs[0], size=size, bold=True, color=accent)


# ---------- 模板辅助 ----------

def _clear_slides(prs: Any) -> None:
    """删除模板自带的示例幻灯片（保留母版/版式/主题），避免成品混入占位内容。"""
    from pptx.oxml.ns import qn

    sld_id_lst = prs.slides._sldIdLst  # noqa: SLF001 - python-pptx 无公开删除 API
    for sld_id in list(sld_id_lst):
        r_id = sld_id.get(qn("r:id"))
        try:
            prs.part.drop_rel(r_id)
        except Exception:  # noqa: BLE001 - 关系缺失不应阻断渲染
            pass
        sld_id_lst.remove(sld_id)


def _layout_types(layout: Any) -> set:
    out: set = set()
    for ph in layout.placeholders:
        try:
            out.add(ph.placeholder_format.type)
        except Exception:  # noqa: BLE001
            continue
    return out


def _pick_layouts(prs: Any) -> tuple[Any, Any]:
    """挑出（标题版式, 内容版式）。按名称与占位符组合启发式匹配，全部有兜底。"""
    from pptx.enum.shapes import PP_PLACEHOLDER

    layouts = [lay for master in prs.slide_masters for lay in master.slide_layouts]
    if not layouts:
        layouts = list(prs.slide_layouts)
    title_types = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
    body_types = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.VERTICAL_BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.VERTICAL_OBJECT}

    title_lay = content_lay = None
    for lay in layouts:
        types = _layout_types(lay)
        name = (lay.name or "").lower()
        has_title = bool(types & title_types)
        has_body = bool(types & body_types)
        if title_lay is None and ("title slide" in name or (has_title and PP_PLACEHOLDER.SUBTITLE in types)):
            title_lay = lay
        if content_lay is None and ("content" in name or (has_title and has_body)):
            content_lay = lay

    if title_lay is None:
        title_lay = next((l for l in layouts if _layout_types(l) & title_types), layouts[0])
    if content_lay is None:
        content_lay = next((l for l in layouts if _layout_types(l) & body_types), layouts[-1])
    return title_lay, content_lay


def _find_body(slide: Any) -> Any | None:
    from pptx.enum.shapes import PP_PLACEHOLDER

    skip = {
        PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE,
        PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER,
        PP_PLACEHOLDER.HEADER, PP_PLACEHOLDER.SLIDE_NUMBER,
    }
    body_types = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.VERTICAL_BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.VERTICAL_OBJECT}
    fallback = None
    for ph in slide.placeholders:
        try:
            ptype = ph.placeholder_format.type
        except Exception:  # noqa: BLE001
            continue
        if ptype in body_types:
            return ph
        if ptype not in skip and fallback is None:
            fallback = ph
    return fallback


def _find_subtitle(slide: Any) -> Any | None:
    from pptx.enum.shapes import PP_PLACEHOLDER

    for ph in slide.placeholders:
        try:
            if ph.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE:
                return ph
        except Exception:  # noqa: BLE001
            continue
    phs = list(slide.placeholders)
    return phs[1] if len(phs) > 1 else None


def _add_textbox(prs: Any, slide: Any, *, top_ratio: float, height_ratio: float) -> Any:
    left = int(prs.slide_width * 0.08)
    top = int(prs.slide_height * top_ratio)
    width = int(prs.slide_width * 0.84)
    height = int(prs.slide_height * height_ratio)
    return slide.shapes.add_textbox(left, top, width, height)


def _write_bullets(tf: Any, bullets: list[str], accent: tuple[int, int, int] | None) -> None:
    tf.clear()
    for i, line in enumerate(bullets[:_MAX_BULLETS]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        if accent is not None:
            for r in p.runs:
                _set_font(r, size=18)


def _set_title(prs: Any, slide: Any, text: str, accent: tuple[int, int, int] | None, size: int) -> None:
    ph = slide.shapes.title
    if ph is None:
        box = _add_textbox(prs, slide, top_ratio=0.06, height_ratio=0.16)
        box.text_frame.text = text
        ph = box
    else:
        ph.text_frame.text = text
    if accent is not None:
        _style_title(ph, accent, size)


def _set_subtitle(prs: Any, slide: Any, text: str, accent: tuple[int, int, int] | None) -> None:
    ph = _find_subtitle(slide)
    if ph is None:
        return
    ph.text_frame.text = text
    if accent is not None:
        for p in ph.text_frame.paragraphs:
            for r in p.runs:
                _set_font(r, size=16)


def _set_body(prs: Any, slide: Any, bullets: list[str], accent: tuple[int, int, int] | None) -> None:
    if not bullets:
        return
    ph = _find_body(slide)
    tf = ph.text_frame if ph is not None else _add_textbox(prs, slide, top_ratio=0.30, height_ratio=0.58).text_frame
    _write_bullets(tf, bullets, accent)


def _cover_subtitle(outline: dict[str, Any]) -> str:
    subtitle = str(outline.get("subtitle") or "").strip()
    stamp = datetime.now(_BEIJING).strftime("%Y-%m-%d")
    return f"{subtitle}\n{stamp}" if subtitle else stamp


def render_pptx(outline: dict[str, Any], style: str = "简洁商务", template_bytes: bytes | None = None) -> bytes:
    """大纲 dict → .pptx 字节。sections 为空也不会抛错（只出封面+封底）。

    template_bytes 提供时套用该模板的母版/主题/背景/Logo；否则用默认版式并按 style 上色。
    """
    from pptx import Presentation

    if template_bytes:
        prs = Presentation(io.BytesIO(template_bytes))
        _clear_slides(prs)
        title_layout, content_layout = _pick_layouts(prs)
        accent = None  # 模板模式沿用主题配色/字体
    else:
        prs = Presentation()
        title_layout = prs.slide_layouts[0]      # Title Slide
        content_layout = prs.slide_layouts[1]    # Title and Content
        accent = _rgb(_STYLE_ACCENT.get(style, _DEFAULT_ACCENT))

    # ---- 封面 ----
    cover = prs.slides.add_slide(title_layout)
    _set_title(prs, cover, str(outline.get("title") or "未命名演示文稿"), accent, 40)
    _set_subtitle(prs, cover, _cover_subtitle(outline), accent)

    # ---- 内容页 ----
    for sec in outline.get("sections") or []:
        heading = str(sec.get("heading") or "").strip()
        bullets = [str(b).strip() for b in (sec.get("bullets") or []) if str(b).strip()]
        if not heading and not bullets:
            continue
        slide = prs.slides.add_slide(content_layout)
        _set_title(prs, slide, heading or " ", accent, 28)
        _set_body(prs, slide, bullets, accent)

    # ---- 封底 ----
    closing = str(outline.get("closing") or "谢谢观看").strip()
    end = prs.slides.add_slide(title_layout)
    _set_title(prs, end, closing, accent, 36)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
